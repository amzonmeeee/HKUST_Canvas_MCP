from __future__ import annotations

import asyncio
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from unittest.mock import Mock
from uuid import uuid4

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfWriter

from webapp.api import create_app
from webapp.db import SCHEMA_VERSION, WorkspaceRepository
from webapp.parsers import (
    ParseError,
    chunks_for_sections,
    extract,
    html_sections,
    parse_file,
)
from webapp.providers import (
    Capabilities,
    HTTPProvider,
    ProviderError,
    validate_base_url,
)
from webapp.services.canvas import CanvasService, CanvasServiceError
from webapp.services.study import artifact_schema
from webapp.store import StudyStore

ORIGIN = "http://127.0.0.1:8765"


class MemorySecrets:
    def __init__(self):
        self.values = {}

    def get(self, identifier):
        return self.values.get(identifier)

    def set(self, identifier, secret):
        self.values[identifier] = secret

    def delete(self, identifier):
        self.values.pop(identifier, None)


class FakeProvider:
    capabilities = Capabilities(native_tools=True)

    def __init__(self):
        self.calls = []
        self.events = None
        self.structured = None
        self.failure = None

    async def stream_chat(self, messages, tools=None):
        self.calls.append((messages, tools))
        if self.events is not None:
            for event in self.events:
                yield event
            return
        context = next(
            (
                m["content"]
                for m in messages
                if m["role"] == "system" and m["content"].startswith("Retrieved source")
            ),
            None,
        )
        chunks = json.loads(context.split("\n", 1)[1]) if context else []
        yield {"type": "text_delta", "text": "Synthetic answer."}
        if chunks:
            yield {
                "type": "text_delta",
                "text": " [cite:" + chunks[0]["chunk_id"] + "]",
            }
        if self.failure:
            raise self.failure
        yield {"type": "message_end"}

    async def generate_structured(self, messages, schema):
        self.calls.append((messages, schema))
        if self.structured is not None:
            return self.structured
        context = next(
            m["content"]
            for m in messages
            if m["content"].startswith("Retrieved source")
        )
        chunk_id = json.loads(context.split("\n", 1)[1])[0]["chunk_id"]
        key = next(k for k in schema["properties"] if k != "title")
        count = schema["properties"][key]["minItems"]
        item = {"citations": [chunk_id]}
        if key == "questions":
            item.update(
                type="multiple_choice",
                question="What is retrieval?",
                choices=["Evidence search", "A browser cookie", "A grade", "An upload"],
                answer_index=0,
                explanation="Retrieval searches evidence.",
            )
        elif key == "cards":
            item.update(front="Retrieval", back="Searching selected evidence.")
        else:
            item.update(
                heading="Retrieval", body="Use selected evidence to support answers."
            )
        return {
            "title": "Synthetic study material",
            key: [dict(item) for _ in range(count)],
        }


@pytest.fixture
def setup(tmp_path, monkeypatch):
    # No browser, Canvas, keychain, or paid provider calls in this suite.
    monkeypatch.setenv("CANVASMCP_CONFIRMATION_DIR", str(tmp_path / "confirmations"))
    canvas = Mock(spec=CanvasService)
    canvas.profile_name.return_value = "Synthetic profile"
    canvas.course.return_value = {
        "id": "101",
        "name": "Synthetic course",
        "course_code": "TEST1000",
    }
    fake, secrets = FakeProvider(), MemorySecrets()
    app = create_app(
        data_dir=tmp_path / "private-data",
        launch_secret="test-launch",
        canvas_service=canvas,
        secret_store=secrets,
        provider_factory=lambda config: fake,
    )
    with TestClient(app, base_url=ORIGIN) as client:
        session = client.post(
            "/api/session",
            headers={"Origin": ORIGIN, "Authorization": "Bearer test-launch"},
        )
        client.headers.update(
            {"Origin": ORIGIN, "X-Workbench-CSRF": session.json()["csrf_token"]}
        )
        workspace = client.post(
            "/api/workspaces", json={"title": "Synthetic readings"}
        ).json()
        provider = client.post(
            "/api/providers",
            json={
                "name": "Synthetic local",
                "kind": "compatible",
                "model": "test-model",
                "base_url": "http://127.0.0.1:11434/v1",
                "native_tools": True,
            },
        ).json()
        yield client, app, workspace, provider, fake, secrets, canvas


def upload_text(
    client,
    workspace,
    title="Retrieval notes",
    content="Retrieval finds relevant evidence in selected course materials.",
):
    result = client.post(
        f"/api/workspaces/{workspace['id']}/sources/text",
        json={"title": title, "content": content},
    )
    assert result.status_code == 201, result.text
    return result.json()


def stream_events(response):
    assert response.status_code == 200, response.text
    return [
        json.loads(block[6:])
        for block in response.text.split("\n\n")
        if block.startswith("data: ")
    ]


def test_migration_preserves_v1_workspace(tmp_path):
    root = tmp_path / "v1"
    root.mkdir()
    identifier = str(uuid4())
    with sqlite3.connect(root / "app.db") as db:
        db.execute(
            "CREATE TABLE workspaces(id TEXT PRIMARY KEY,kind TEXT,canvas_course_id TEXT UNIQUE,title TEXT,description TEXT,course_code TEXT,term_name TEXT,created_at TEXT,updated_at TEXT,last_sync_at TEXT)"
        )
        db.execute(
            "INSERT INTO workspaces VALUES(?,?,?,?,?,?,?,?,?,?)",
            (
                identifier,
                "custom",
                None,
                "Keep me",
                "Private",
                None,
                None,
                "old",
                "old",
                None,
            ),
        )
        db.execute("PRAGMA user_version=1")
    repo = WorkspaceRepository(root)
    repo.initialize()
    repo.initialize()
    assert repo.get(identifier)["description"] == "Private"
    assert StudyStore(repo).sources(identifier) == []
    with sqlite3.connect(root / "app.db") as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION


def test_upload_parse_search_scope_and_delete(setup):
    client, app, workspace, *_ = setup
    source = upload_text(client, workspace)
    assert source["status"] == "ready"
    assert app.state.study_store.search(workspace["id"], [source["id"]], "retrieval")
    assert app.state.study_store.search(workspace["id"], [], "retrieval") == []
    details = client.get(
        f"/api/workspaces/{workspace['id']}/sources/{source['id']}"
    ).json()
    chunk = details["chunks"][0]
    assert chunk["locator"]["offset"] == 0
    assert (
        client.get(f"/api/workspaces/{workspace['id']}/citations/{chunk['id']}")
        .json()["text"]
        .startswith("Retrieval")
    )
    other = client.post("/api/workspaces", json={"title": "Other"}).json()
    assert (
        client.get(f"/api/workspaces/{other['id']}/sources/{source['id']}").status_code
        == 404
    )
    assert (
        client.get(f"/api/workspaces/{other['id']}/citations/{chunk['id']}").status_code
        == 404
    )
    assert (
        client.delete(
            f"/api/workspaces/{workspace['id']}/sources/{source['id']}"
        ).status_code
        == 204
    )
    assert not (app.state.repository.data_dir / "sources" / source["id"]).exists()
    assert not app.state.study_store.search(
        workspace["id"], [source["id"]], "retrieval"
    )


def test_multipart_upload_private_path_and_parser_failure_isolation(setup):
    client, app, workspace, *_ = setup
    base = f"/api/workspaces/{workspace['id']}/sources"
    uploaded = client.post(
        base + "/upload",
        files={"file": ("../../private.txt", b"Synthetic evidence", "text/plain")},
    )
    assert uploaded.status_code == 201
    source = uploaded.json()
    assert source["title"] == "private.txt" and source["status"] == "ready"
    path = app.state.repository.data_dir / "sources" / source["id"] / "original.txt"
    assert path.stat().st_mode & 0o777 == 0o600
    broken = client.post(
        base + "/upload",
        files={"file": ("malformed.pdf", b"not a pdf", "application/pdf")},
    ).json()
    assert broken["status"] == "failed" and broken["error_code"] == "parser_failed"
    unsupported = client.post(
        base + "/upload",
        files={"file": ("recording.mp4", b"not downloaded", "video/mp4")},
    ).json()
    assert unsupported["status"] == "unsupported"
    assert len(app.state.study_store.sources(workspace["id"])) == 3
    assert app.state.study_store.search(workspace["id"], [source["id"]], "evidence")


def test_workspace_delete_removes_fts_and_files(setup):
    client, app, workspace, *_ = setup
    source = upload_text(client, workspace)
    assert client.delete(f"/api/workspaces/{workspace['id']}").status_code == 204
    assert not (app.state.repository.data_dir / "sources" / source["id"]).exists()
    with sqlite3.connect(app.state.repository.path) as db:
        assert db.execute("SELECT count(*) FROM chunks_fts").fetchone()[0] == 0


def test_html_heading_excludes_scripts_and_overlap():
    chunks = chunks_for_sections(
        html_sections(
            "<h1>Evidence</h1><p>Safe text</p><script>steal secrets</script><style>private</style><h2>Details</h2><p>More evidence</p>"
        )
    )
    assert chunks[0]["locator"]["heading"] == "Evidence"
    assert chunks[1]["locator"]["heading"] == "Details"
    assert "steal" not in str(chunks) and "private" not in str(chunks)
    long = chunks_for_sections([("word " * 1000, {"page": 7})])
    assert len(long) > 2 and all(c["locator"]["page"] == 7 for c in long)
    assert all(len(c["text"]) <= 1800 for c in long)


@pytest.mark.parametrize(
    "suffix,text,locator",
    [
        (".md", "# Evidence\nRetrieval is useful.", "heading"),
        (
            ".vtt",
            "WEBVTT\n\n00:00:02.000 --> 00:00:04.000\nRetrieval evidence.",
            "timestamp",
        ),
        (".srt", "1\n00:00:02,000 --> 00:00:04,000\nRetrieval evidence.", "timestamp"),
        (".html", "<h2>Retrieval</h2><p>Evidence</p>", "heading"),
    ],
)
def test_text_formats_have_locators(tmp_path, suffix, text, locator):
    path = tmp_path / ("source" + suffix)
    path.write_text(text)
    assert extract(path, suffix)[0]["locator"][locator]


def test_document_formats_and_empty_pdf(tmp_path):
    doc = Document()
    doc.add_heading("Retrieval", level=1)
    doc.add_paragraph("Search selected evidence.")
    path = tmp_path / "source.docx"
    doc.save(path)
    chunks = parse_file(path, ".docx")
    assert any("Search selected" in c["text"] for c in chunks)
    assert chunks[1]["locator"]["paragraph"] == 2
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Retrieval"
    slide.placeholders[1].text = "Evidence selection"
    slides = tmp_path / "slides.pptx"
    presentation.save(slides)
    assert parse_file(slides, ".pptx")[0]["locator"]["slide"] == 1
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf = tmp_path / "scan.pdf"
    writer.write(pdf)
    with pytest.raises(ParseError, match="No readable text"):
        parse_file(pdf, ".pdf")


def test_cjk_search_and_malicious_fts_query(setup):
    client, app, workspace, *_ = setup
    source = upload_text(client, workspace, content="這裡講解機器學習及檢索方法。")
    assert app.state.study_store.search(workspace["id"], [source["id"]], "機器學習")
    assert (
        app.state.study_store.search(
            workspace["id"], [source["id"]], '" OR * NEAR(abc)'
        )
        == []
    )


def test_incremental_sync_reuses_chunks_and_reports_independent_failures(
    setup, tmp_path
):
    client, app, _workspace, _provider, _fake, _secrets, canvas = setup
    course = client.post(
        "/api/workspaces", json={"kind": "canvas_course", "canvas_course_id": "101"}
    ).json()

    def call(method, **args):
        if method == "get_course":
            return {
                "syllabus_body": "<h1>Retrieval</h1><p>Evidence search.</p>",
                "updated_at": "v1",
            }
        if method == "list_pages":
            return [
                {"url": "retrieval", "title": "Retrieval page", "updated_at": "v1"},
                {"url": "denied", "title": "Unavailable", "updated_at": "v1"},
            ]
        if method == "get_page":
            if args["url_or_id"] == "denied":
                raise CanvasServiceError(
                    "canvas_permission_denied", "Access denied.", 403
                )
            return {
                "body": "<h2>Retrieval</h2><p>Evidence search.</p>",
                "updated_at": "v1",
            }
        if method == "activity_list":
            return {
                "items": [
                    {
                        "id": 10,
                        "type": "ExternalUrl",
                        "title": "External video",
                        "external_url": "https://example.com/video?token=secret",
                    }
                ],
                "truncated": False,
            }
        if method == "list_modules":
            return [
                {
                    "id": 1,
                    "name": "Week 1",
                    "items": [
                        {
                            "id": 10,
                            "type": "ExternalUrl",
                            "title": "External video",
                            "external_url": "https://example.com/video?token=secret",
                        }
                    ],
                }
            ]
        return []

    canvas.client_call.side_effect = call
    result = client.post(f"/api/workspaces/{course['id']}/sources/inventory").json()
    assert any(s["status"] == "external_reference" for s in result["sources"])
    assert "secret" not in json.dumps(result)
    identifiers = [s["id"] for s in result["sources"]]
    synced = client.post(
        f"/api/workspaces/{course['id']}/sources/sync", json={"source_ids": identifiers}
    ).json()["sources"]
    ready = next(s for s in synced if s["title"] == "Retrieval page")
    assert any(s["status"] == "failed" for s in synced)
    ids_before = [
        c["id"] for c in app.state.study_store.chunks(course["id"], ready["id"])
    ]
    client.post(
        f"/api/workspaces/{course['id']}/sources/sync",
        json={"source_ids": [ready["id"]]},
    )
    assert [
        c["id"] for c in app.state.study_store.chunks(course["id"], ready["id"])
    ] == ids_before
    # Ordinary reads never contact Canvas.
    canvas.client_call.reset_mock()
    client.get(f"/api/workspaces/{course['id']}/sources")
    canvas.client_call.assert_not_called()


def test_provider_secrets_never_echo_or_persist(setup):
    client, app, _workspace, _provider, fake, secrets, _canvas = setup
    key = "synthetic-api-key-do-not-echo"
    config = {
        "name": "Official test",
        "kind": "openai",
        "model": "test-model",
        "api_key": key,
    }
    saved = client.post("/api/providers", json=config)
    assert saved.status_code == 201
    identifier = saved.json()["id"]
    assert secrets.get(identifier) == key
    assert key not in saved.text and key not in client.get("/api/providers").text
    assert key.encode() not in app.state.repository.path.read_bytes()
    invalid = client.post("/api/providers", json={**config, "api_key": key + "\n"})
    assert invalid.status_code == 422 and key not in invalid.text
    assert (
        client.post(f"/api/providers/{identifier}/test").json()["status"] == "connected"
    )
    assert "course" not in fake.calls[-1][0][0]["content"].split(";")[0]
    assert client.delete(f"/api/providers/{identifier}").status_code == 204
    assert secrets.get(identifier) is None


@pytest.mark.parametrize(
    "url",
    [
        "http://example.com/v1",
        "https://example.com/v1?token=x",
        "https://user:pass@example.com/v1",
        "file:///tmp/v1",
        "https://example.com/api",
    ],
)
def test_provider_urls_reject_unsafe_configuration(url):
    with pytest.raises(ProviderError):
        validate_base_url("compatible", url)


def test_endpoint_change_requires_removing_saved_key(setup):
    client, app, _workspace, provider, _fake, secrets, _canvas = setup
    secrets.set(provider["id"], "synthetic-key")
    result = client.post(
        "/api/providers",
        json={
            "id": provider["id"],
            "name": "Changed",
            "kind": "compatible",
            "model": "test",
            "base_url": "https://example.com/v1",
        },
    )
    assert result.status_code == 422
    assert (
        app.state.study_store.provider(provider["id"])["base_url"]
        == "http://127.0.0.1:11434/v1"
    )


def test_grounded_chat_stream_persistence_citations_and_notes(setup):
    client, _app, workspace, provider, fake, *_ = setup
    source = upload_text(client, workspace)
    upload_text(
        client, workspace, title="Excluded", content="PRIVATE UNSELECTED SENTINEL"
    )
    base = f"/api/workspaces/{workspace['id']}"
    events = stream_events(
        client.post(
            base + "/chat",
            json={
                "provider_id": provider["id"],
                "source_ids": [source["id"]],
                "message": "Explain retrieval",
            },
        )
    )
    assert any(e["type"] == "text_delta" for e in events)
    citation_event = next(e for e in events if e["type"] == "citation")
    assert citation_event["citation"]["source_id"] == source["id"]
    assert "UNSELECTED SENTINEL" not in json.dumps(fake.calls)
    conversation_id = next(
        e["conversation_id"] for e in events if e["type"] == "conversation"
    )
    saved = client.get(base + "/conversations/" + conversation_id).json()
    assert len(saved["messages"]) == 2 and saved["messages"][1]["status"] == "complete"
    note = client.post(
        base + "/notes",
        json={
            "title": "Saved answer",
            "content": "",
            "message_id": saved["messages"][1]["id"],
        },
    ).json()
    assert note["provenance"]["citations"][0]["source_id"] == source["id"]
    changed = client.put(
        base + "/notes/" + note["id"],
        json={"title": "Edited", "content": "My local notes"},
    ).json()
    assert (
        changed["content"] == "My local notes"
        and changed["provenance"] == note["provenance"]
    )
    assert client.delete(base + "/notes/" + note["id"]).status_code == 204
    assert client.delete(base + "/conversations/" + conversation_id).status_code == 204


def test_chat_error_saves_partial_and_invalid_citation_is_removed(setup):
    client, app, workspace, provider, fake, *_ = setup
    fake.events = [
        {"type": "text_delta", "text": "Partial [cite:invented]"},
        {"type": "message_end"},
    ]
    base = f"/api/workspaces/{workspace['id']}"
    events = stream_events(
        client.post(
            base + "/chat",
            json={
                "provider_id": provider["id"],
                "message": "General knowledge",
                "source_ids": [],
            },
        )
    )
    assert any(e["type"] == "warning" for e in events)
    assert "[cite:invented]" not in events[-1]["content"]
    fake.events = None
    fake.failure = ProviderError("provider_rate_limit", "Quota exhausted.")
    events = stream_events(
        client.post(
            base + "/chat", json={"provider_id": provider["id"], "message": "Try again"}
        )
    )
    assert events[-1]["type"] == "error"
    identifier = events[0]["conversation_id"]
    saved = client.get(base + "/conversations/" + identifier).json()["messages"][-1]
    assert saved["content"] == "Synthetic answer." and saved["status"] == "failed"
    assert not app.state.study._running


def test_chat_rejects_foreign_sources_and_unsupported_tools(setup):
    client, _app, workspace, provider, fake, *_ = setup
    other = client.post("/api/workspaces", json={"title": "Other"}).json()
    source = upload_text(client, other)
    payload = {
        "provider_id": provider["id"],
        "message": "Evidence",
        "source_ids": [source["id"]],
    }
    base = f"/api/workspaces/{workspace['id']}"
    assert client.post(base + "/chat", json=payload).status_code == 422
    fake.capabilities = Capabilities(native_tools=False)
    assert (
        client.post(
            base + "/chat", json={**payload, "source_ids": [], "live_tools": True}
        ).status_code
        == 422
    )
    assert not fake.calls


@pytest.mark.parametrize("kind", ["quiz", "flashcards", "study_guide"])
def test_artifact_contract_provenance_exports_and_saved_notes(setup, kind):
    client, _app, workspace, provider, _fake, *_ = setup
    source = upload_text(client, workspace)
    base = f"/api/workspaces/{workspace['id']}"
    result = client.post(
        base + "/artifacts/generate",
        json={
            "provider_id": provider["id"],
            "source_ids": [source["id"]],
            "kind": kind,
            "count": 2,
            "difficulty": "advanced",
        },
    )
    assert result.status_code == 201, result.text
    artifact = result.json()
    assert artifact["provenance"]["source_ids"] == [source["id"]]
    assert artifact["provenance"]["model"] == "test-model"
    assert len(artifact["provenance"]["citations"]) == 1
    exported = client.get(base + f"/artifacts/{artifact['id']}/export?format=markdown")
    assert exported.status_code == 200 and "## Provenance" in exported.text
    assert "attachment" in exported.headers["content-disposition"]
    assert (
        client.get(base + f"/artifacts/{artifact['id']}/export").json()["id"]
        == artifact["id"]
    )
    note = client.post(
        base + "/notes",
        json={"title": "Saved material", "content": "", "artifact_id": artifact["id"]},
    ).json()
    assert note["provenance"]["artifact_id"] == artifact["id"]
    assert client.delete(base + f"/artifacts/{artifact['id']}").status_code == 204
    assert (
        client.get(base + "/notes")
        .json()["notes"][0]["content"]
        .startswith("# Synthetic")
    )


def test_artifacts_refuse_hallucinated_citations(setup):
    client, _app, workspace, provider, fake, *_ = setup
    source = upload_text(client, workspace)
    fake.structured = {
        "title": "Bad citations",
        "cards": [{"front": "Q", "back": "A", "citations": [str(uuid4())]}],
    }
    base = f"/api/workspaces/{workspace['id']}"
    response = client.post(
        base + "/artifacts/generate",
        json={
            "provider_id": provider["id"],
            "source_ids": [source["id"]],
            "kind": "flashcards",
            "count": 1,
        },
    )
    assert (
        response.status_code == 422
        and response.json()["error"]["code"] == "invalid_citations"
    )
    assert client.get(base + "/artifacts").json()["artifacts"] == []


def test_writes_preview_only_exact_confirm_and_single_use(setup):
    client, _app, _workspace, _provider, _fake, _secrets, canvas = setup
    course = client.post(
        "/api/workspaces", json={"kind": "canvas_course", "canvas_course_id": "101"}
    ).json()
    canvas._invoke.return_value = {
        "status": "preview",
        "written": False,
        "account": {"user_id": "1", "name": "Synthetic user"},
        "target": {"course_id": "101", "topic_id": "12"},
        "payload": {"message": "Synthetic preview"},
        "confirmation_token": "private-confirmation-token",
        "expires_at": "later",
    }
    base = f"/api/workspaces/{course['id']}/actions"
    args = {"course_id": "101", "topic_id": "12", "message": "Synthetic preview"}
    response = client.post(
        base + "/preview", json={"name": "post_discussion_entry", "arguments": args}
    )
    assert (
        response.status_code == 200
        and "private-confirmation-token" not in response.text
    )
    canvas._invoke.assert_called_once_with("post_discussion_entry", args)
    preview = response.json()["preview"]
    assert (
        client.post(
            base + f"/{preview['id']}/confirm",
            json={"approved": True, "payload": {"message": "tampered"}},
        ).status_code
        == 422
    )
    canvas._invoke.return_value = {"status": "written", "written": True}
    assert (
        client.post(
            base + f"/{preview['id']}/confirm", json={"approved": True}
        ).status_code
        == 200
    )
    canvas._invoke.assert_called_with(
        "post_discussion_entry",
        {**args, "confirmation_token": "private-confirmation-token"},
    )
    assert (
        client.post(
            base + f"/{preview['id']}/confirm", json={"approved": True}
        ).status_code
        == 409
    )
    assert canvas._invoke.call_count == 2
    assert (
        client.post(
            base + "/preview",
            json={
                "name": "post_discussion_entry",
                "arguments": {**args, "confirmation_token": "forged"},
            },
        ).status_code
        == 422
    )
    assert (
        client.post(
            base + "/preview",
            json={
                "name": "post_discussion_entry",
                "arguments": {**args, "course_id": "999"},
            },
        ).status_code
        == 422
    )


def test_model_write_is_a_preview_and_never_confirmed(setup):
    client, _app, workspace, provider, fake, _secrets, canvas = setup
    canvas._invoke.return_value = {
        "status": "preview",
        "written": False,
        "account": {"user_id": "1"},
        "target": {"conversation_id": "3"},
        "payload": {"body": "Synthetic reply"},
        "confirmation_token": "server-only-token",
    }
    events = [
        {
            "type": "tool_call_end",
            "id": "call1",
            "name": "reply_to_conversation",
            "arguments": {"conversation_id": "3", "body": "Synthetic reply"},
        },
        {"type": "message_end"},
    ]

    async def stream(messages, tools=None):
        fake.calls.append((messages, tools))
        for event in (
            events
            if len(fake.calls) == 1
            else [
                {"type": "text_delta", "text": "Please review the preview."},
                {"type": "message_end"},
            ]
        ):
            yield event

    fake.stream_chat = stream
    result = stream_events(
        client.post(
            f"/api/workspaces/{workspace['id']}/chat",
            json={
                "provider_id": provider["id"],
                "message": "Prepare a reply",
                "live_tools": True,
            },
        )
    )
    assert any(e["type"] == "write_preview" for e in result)
    assert canvas._invoke.call_count == 1
    assert "server-only-token" not in json.dumps(
        fake.calls
    ) and "server-only-token" not in json.dumps(result)


def test_cancelled_preview_cannot_write_and_capabilities_hide_confirmation(setup):
    client, _app, workspace, _provider, _fake, _secrets, canvas = setup
    canvas._invoke.return_value = {
        "status": "preview",
        "written": False,
        "confirmation_token": "private-token",
    }
    base = f"/api/workspaces/{workspace['id']}/actions"
    tools = client.get(base).json()["tools"]
    assert all("confirmation_token" not in t["parameters"]["properties"] for t in tools)
    assert not any("submit" in t["name"] or "download" in t["name"] for t in tools)
    preview = client.post(
        base + "/preview",
        json={
            "name": "reply_to_conversation",
            "arguments": {"conversation_id": "3", "body": "Test"},
        },
    ).json()["preview"]
    assert client.delete(base + "/" + preview["id"]).status_code == 204
    assert (
        client.post(
            base + f"/{preview['id']}/confirm", json={"approved": True}
        ).status_code
        == 409
    )
    assert canvas._invoke.call_count == 1


def config(kind, streaming=True):
    return {
        "id": str(uuid4()),
        "name": "Synthetic",
        "kind": kind,
        "model": "test-model",
        "base_url": "https://example.com/v1",
        "capabilities": asdict(Capabilities(streaming=streaming, native_tools=True)),
    }


def collect(provider, messages=None, tools=None):
    async def run():
        return [
            e
            async for e in provider.stream_chat(
                messages or [{"role": "user", "content": "Test"}], tools
            )
        ]

    return asyncio.run(run())


@pytest.mark.parametrize(
    "kind,wire",
    [
        (
            "openai",
            [
                {"type": "response.output_text.delta", "delta": "你好"},
                {"type": "response.completed"},
            ],
        ),
        (
            "anthropic",
            [
                {
                    "type": "content_block_delta",
                    "index": 0,
                    "delta": {"type": "text_delta", "text": "你好"},
                },
                {"type": "message_stop"},
            ],
        ),
        (
            "compatible",
            [
                {"choices": [{"delta": {"content": "你好"}, "finish_reason": None}]},
                {"choices": [{"delta": {}, "finish_reason": "stop"}]},
            ],
        ),
    ],
)
def test_provider_stream_contract_and_wire_payload(kind, wire):
    requests = []

    def handle(request):
        requests.append(request)
        return httpx.Response(
            200,
            text="".join("data: " + json.dumps(e) + "\n\n" for e in wire),
            headers={"Content-Type": "text/event-stream"},
        )

    provider = HTTPProvider(
        config(kind), "synthetic-key", transport=httpx.MockTransport(handle)
    )
    events = collect(provider)
    assert events[0]["type"] == "message_start" and events[-1]["type"] == "message_end"
    assert next(e for e in events if e["type"] == "text_delta")["text"] == "你好"
    body = json.loads(requests[0].content)
    assert body["stream"] is True and body["model"] == "test-model"
    if kind == "openai":
        assert body["store"] is False and requests[0].url.path == "/v1/responses"
    elif kind == "anthropic":
        assert requests[0].headers["anthropic-version"] == "2023-06-01"
    else:
        assert requests[0].url.path == "/v1/chat/completions"


@pytest.mark.parametrize(
    "kind,wire",
    [
        (
            "openai",
            [
                {
                    "type": "response.output_item.added",
                    "output_index": 0,
                    "item": {
                        "type": "function_call",
                        "call_id": "c1",
                        "name": "test",
                        "arguments": "",
                    },
                },
                {
                    "type": "response.function_call_arguments.delta",
                    "output_index": 0,
                    "delta": '{"id":',
                },
                {
                    "type": "response.function_call_arguments.done",
                    "output_index": 0,
                    "arguments": '{"id":"1"}',
                },
                {"type": "response.completed"},
            ],
        ),
        (
            "anthropic",
            [
                {
                    "type": "content_block_start",
                    "index": 1,
                    "content_block": {
                        "type": "tool_use",
                        "id": "c1",
                        "name": "test",
                        "input": {},
                    },
                },
                {
                    "type": "content_block_delta",
                    "index": 1,
                    "delta": {"type": "input_json_delta", "partial_json": '{"id":"1"}'},
                },
                {"type": "message_stop"},
            ],
        ),
        (
            "compatible",
            [
                {
                    "choices": [
                        {
                            "delta": {
                                "tool_calls": [
                                    {
                                        "index": 0,
                                        "id": "c1",
                                        "function": {
                                            "name": "test",
                                            "arguments": '{"id":',
                                        },
                                    }
                                ]
                            }
                        }
                    ]
                },
                {
                    "choices": [
                        {
                            "delta": {
                                "tool_calls": [
                                    {"index": 0, "function": {"arguments": '"1"}'}}
                                ]
                            },
                            "finish_reason": "tool_calls",
                        }
                    ]
                },
            ],
        ),
    ],
)
def test_native_tool_stream_contract(kind, wire):
    transport = httpx.MockTransport(
        lambda request: httpx.Response(
            200, text="".join("data: " + json.dumps(e) + "\n\n" for e in wire)
        )
    )
    events = collect(
        HTTPProvider(config(kind), "synthetic-key", transport=transport),
        tools=[
            {
                "name": "test",
                "description": "Test",
                "parameters": {
                    "type": "object",
                    "properties": {"id": {"type": "string"}},
                },
            }
        ],
    )
    call = next(e for e in events if e["type"] == "tool_call_end")
    assert call == {
        "type": "tool_call_end",
        "id": "c1",
        "name": "test",
        "arguments": {"id": "1"},
    }


@pytest.mark.parametrize(
    "status,code",
    [
        (401, "provider_auth"),
        (403, "provider_auth"),
        (429, "provider_rate_limit"),
        (500, "provider_unavailable"),
        (302, "provider_unavailable"),
    ],
)
def test_provider_errors_are_sanitized_and_redirects_never_followed(status, code):
    requests = []

    def handle(request):
        requests.append(request)
        return httpx.Response(
            status,
            text="PRIVATE RAW ERROR api-key",
            headers={"Location": "https://malicious.example/v1"},
        )

    with pytest.raises(ProviderError) as caught:
        collect(
            HTTPProvider(
                config("openai"), "synthetic-key", transport=httpx.MockTransport(handle)
            )
        )
    assert caught.value.code == code and "PRIVATE" not in str(caught.value)
    assert len(requests) == 1


def test_provider_incomplete_stream_is_not_success():
    transport = httpx.MockTransport(
        lambda request: httpx.Response(
            200,
            text='data: {"type":"response.output_text.delta","delta":"partial"}\n\n',
        )
    )
    with pytest.raises(ProviderError) as caught:
        collect(HTTPProvider(config("openai"), "synthetic-key", transport=transport))
    assert caught.value.code == "provider_stream_incomplete"


@pytest.mark.parametrize(
    "kind,response",
    [
        (
            "openai",
            {
                "status": "completed",
                "output": [
                    {
                        "type": "message",
                        "content": [{"type": "output_text", "text": "OK"}],
                    }
                ],
            },
        ),
        (
            "anthropic",
            {"content": [{"type": "text", "text": "OK"}], "stop_reason": "end_turn"},
        ),
        (
            "compatible",
            {"choices": [{"message": {"content": "OK"}, "finish_reason": "stop"}]},
        ),
    ],
)
def test_nonstreaming_provider_support(kind, response):
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json=response))
    events = collect(
        HTTPProvider(
            config(kind, streaming=False), "synthetic-key", transport=transport
        )
    )
    assert events[1] == {"type": "text_delta", "text": "OK"}
    assert events[-1]["type"] == "message_end"


@pytest.mark.parametrize("kind", ["openai", "anthropic", "compatible"])
def test_native_structured_output_wire_contract(kind):
    requests = []
    schema = artifact_schema("flashcards", 1)
    content = json.dumps(
        {
            "title": "Synthetic cards",
            "cards": [{"front": "Q", "back": "A", "citations": ["synthetic-chunk"]}],
        }
    )

    def handle(request):
        requests.append(json.loads(request.content))
        if kind == "openai":
            wire = [
                {"type": "response.output_text.delta", "delta": content},
                {"type": "response.completed"},
            ]
        elif kind == "anthropic":
            wire = [
                {
                    "type": "content_block_delta",
                    "index": 0,
                    "delta": {"type": "text_delta", "text": content},
                },
                {"type": "message_stop"},
            ]
        else:
            wire = [
                {"choices": [{"delta": {"content": content}, "finish_reason": "stop"}]}
            ]
        return httpx.Response(
            200, text="".join("data: " + json.dumps(event) + "\n\n" for event in wire)
        )

    configuration = config(kind)
    configuration["capabilities"]["structured_output"] = True
    result = asyncio.run(
        HTTPProvider(
            configuration, "synthetic-key", transport=httpx.MockTransport(handle)
        ).generate_structured([{"role": "user", "content": "Create cards"}], schema)
    )
    assert result["title"] == "Synthetic cards"
    body = requests[0]
    native = (
        body["text"]["format"]["schema"]
        if kind == "openai"
        else body["output_config"]["format"]["schema"]
        if kind == "anthropic"
        else body["response_format"]["json_schema"]["schema"]
    )
    assert (
        native["additionalProperties"] is False
        and "minItems" not in native["properties"]["cards"]
    )


def test_structured_fallback_repairs_once_and_does_not_retry_network_failure():
    requests = []

    def handle(request):
        requests.append(json.loads(request.content))
        answer = "invalid JSON" if len(requests) == 1 else '{"answer":"OK"}'
        return httpx.Response(
            200,
            text="data: "
            + json.dumps(
                {"choices": [{"delta": {"content": answer}, "finish_reason": "stop"}]}
            )
            + "\n\n",
        )

    schema = {
        "type": "object",
        "required": ["answer"],
        "properties": {"answer": {"type": "string"}},
        "additionalProperties": False,
    }
    provider = HTTPProvider(config("compatible"), transport=httpx.MockTransport(handle))
    assert asyncio.run(
        provider.generate_structured(
            [{"role": "user", "content": "JSON please"}], schema
        )
    ) == {"answer": "OK"}
    assert (
        len(requests) == 2
        and "Repair your previous answer" in requests[1]["messages"][-1]["content"]
    )
    failures = []

    def rejected(request):
        failures.append(request)
        return httpx.Response(429)

    with pytest.raises(ProviderError):
        asyncio.run(
            HTTPProvider(
                config("compatible"), transport=httpx.MockTransport(rejected)
            ).generate_structured([{"role": "user", "content": "Test"}], schema)
        )
    assert len(failures) == 1


def test_parallel_database_launches_keep_migrations_atomic(tmp_path):
    root = tmp_path / "shared-app"

    def initialize(_):
        repository = WorkspaceRepository(root)
        repository.initialize()
        return repository

    with ThreadPoolExecutor(max_workers=4) as pool:
        repositories = list(pool.map(initialize, range(16)))
    with sqlite3.connect(repositories[0].path) as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION
    assert repositories[0].create(title="Still works")["title"] == "Still works"


def test_keyless_provider_never_touches_keychain(setup):
    client, _app, _workspace, provider, _fake, secrets, _canvas = setup
    secrets.get = Mock(
        side_effect=AssertionError("Do not request Keychain on ordinary page loads")
    )
    secrets.delete = Mock(side_effect=AssertionError("No key exists to delete"))
    assert client.get("/api/providers").json()["providers"][0]["has_key"] is False
    assert client.delete("/api/providers/" + provider["id"]).status_code == 204


@pytest.mark.parametrize(
    "storage_host",
    [
        "bucket.s3.amazonaws.com",
        "files.instructure.com",
        "a1.cluster1.canvas-user-content.com",
        "inst-fs-sin-prod.inscloudgate.net",
    ],
)
def test_bounded_canvas_download_rejects_external_redirect_and_strips_csrf(
    tmp_path, storage_host
):
    from client.content import CanvasContentMixin

    class Adapter(CanvasContentMixin):
        def _call_canvas(self, callback, context):
            return callback(canvas)

        def _item_to_dict(self, value):
            return {"id": 1, "size": 3}

    def response(status, headers=None, content=b"abc"):
        item = Mock()
        item.status_code, item.headers = status, headers or {}
        item.__enter__ = Mock(return_value=item)
        item.__exit__ = Mock(return_value=False)
        item.iter_content.return_value = [content]
        return item

    file = Mock(url="https://canvas.ust.hk/files/1/download")
    session = Mock()
    canvas = Mock()
    canvas.get_course.return_value.get_file.return_value = file
    canvas._Canvas__requester = Mock(_session=session)
    session.get.side_effect = [
        response(302, {"Location": "https://evil.example/steal"})
    ]
    with pytest.raises(ValueError, match="approved storage"):
        Adapter().download_file_bounded(
            course_id="101", file_id="1", destination_path=str(tmp_path / "file")
        )
    assert session.get.call_count == 1
    session.get.reset_mock()
    session.get.side_effect = [
        response(
            302,
            {"Location": f"https://{storage_host}/file?signature=synthetic"},
        ),
        response(200),
    ]
    destination = tmp_path / "file"
    Adapter().download_file_bounded(
        course_id="101", file_id="1", destination_path=str(destination)
    )
    assert destination.read_bytes() == b"abc"
    assert session.get.call_args_list[1].kwargs["headers"]["X-CSRF-Token"] is None
    assert session.get.call_args_list[1].kwargs["allow_redirects"] is False
    session.get.side_effect = [response(200, content=b"too large")]
    with pytest.raises(ValueError, match="download limit"):
        Adapter().download_file_bounded(
            course_id="101", file_id="1", destination_path=str(destination), max_bytes=4
        )


def test_pdf_extracts_page_locator(tmp_path):
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (Synthetic retrieval evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "text.pdf"
    writer.write(path)
    chunks = parse_file(path, ".pdf")
    assert (
        chunks[0]["locator"]["page"] == 1 and "retrieval evidence" in chunks[0]["text"]
    )


def test_metadata_logging_drops_private_bodies_and_credentials(monkeypatch):
    from webapp.observability import event, logger

    log = Mock()
    monkeypatch.setattr(logger, "info", log)
    event(
        "synthetic_operation",
        status="ready",
        count=2,
        body="PRIVATE BODY",
        api_key="PRIVATE KEY",
        cookies="PRIVATE COOKIES",
        csrf_token="PRIVATE CSRF",
    )
    record = log.call_args.args[0]
    assert "PRIVATE" not in record and json.loads(record) == {
        "event": "synthetic_operation",
        "status": "ready",
        "count": 2,
    }


def test_chat_scope_change_keeps_history_local_but_omits_unselected_context(setup):
    client, _app, workspace, provider, fake, *_ = setup
    source = upload_text(
        client, workspace, content="Retrieval scope-sensitive evidence."
    )
    base = f"/api/workspaces/{workspace['id']}"
    first = stream_events(
        client.post(
            base + "/chat",
            json={
                "provider_id": provider["id"],
                "source_ids": [source["id"]],
                "message": "Explain retrieval: PRIVATE HISTORY SENTINEL",
            },
        )
    )
    conversation_id = first[0]["conversation_id"]
    stored = client.get(base + "/conversations/" + conversation_id).json()["messages"]
    assert stored[0]["source_ids"] == [source["id"]]
    second = stream_events(
        client.post(
            base + "/chat",
            json={
                "provider_id": provider["id"],
                "conversation_id": conversation_id,
                "source_ids": [],
                "message": "Now answer with general knowledge only",
            },
        )
    )
    assert second[-1]["type"] == "message_end"
    assert "PRIVATE HISTORY SENTINEL" not in json.dumps(fake.calls[-1])
    assert (
        len(client.get(base + "/conversations/" + conversation_id).json()["messages"])
        == 4
    )


def test_note_from_artifact_excerpt_pins_saved_item_and_provenance(setup):
    client, app, workspace, provider, fake, *_ = setup
    source = upload_text(client, workspace)
    second_source = upload_text(client, workspace, title="Second evidence")
    chunk_id = app.state.study_store.chunks(workspace["id"], source["id"])[0]["id"]
    second_chunk_id = app.state.study_store.chunks(
        workspace["id"], second_source["id"]
    )[0]["id"]
    fake.structured = {
        "title": "Two cards",
        "cards": [
            {"front": "First", "back": "First answer", "citations": [chunk_id]},
            {
                "front": "Second",
                "back": "Second answer",
                "citations": [second_chunk_id],
            },
        ],
    }
    base = f"/api/workspaces/{workspace['id']}"
    artifact = client.post(
        base + "/artifacts/generate",
        json={
            "provider_id": provider["id"],
            "source_ids": [source["id"], second_source["id"]],
            "kind": "flashcards",
            "count": 2,
        },
    ).json()
    result = client.post(
        base + "/notes",
        json={
            "title": "Selected card",
            "content": "forged content ignored",
            "artifact_id": artifact["id"],
            "artifact_index": 1,
        },
    )
    assert result.status_code == 201
    note = result.json()
    assert (
        "Second answer" in note["content"]
        and "First answer" not in note["content"]
        and "forged" not in note["content"]
    )
    assert note["provenance"]["artifact_index"] == 1
    assert note["provenance"]["source_ids"] == [second_source["id"]]
    assert [c["chunk_id"] for c in note["provenance"]["citations"]] == [second_chunk_id]
    assert (
        client.post(
            base + "/notes",
            json={
                "title": "Bad",
                "content": "",
                "artifact_id": artifact["id"],
                "artifact_index": 5,
            },
        ).status_code
        == 422
    )


def test_provider_context_and_network_errors_are_distinct():
    context_transport = httpx.MockTransport(
        lambda request: httpx.Response(
            400,
            json={
                "error": {
                    "code": "context_length_exceeded",
                    "message": "PRIVATE RAW PROMPT",
                }
            },
        )
    )
    with pytest.raises(ProviderError) as context:
        collect(HTTPProvider(config("compatible"), transport=context_transport))
    assert context.value.code == "provider_context_exceeded" and "PRIVATE" not in str(
        context.value
    )

    def disconnected(request):
        raise httpx.ConnectError("PRIVATE CONNECTION DETAIL")

    with pytest.raises(ProviderError) as network:
        collect(
            HTTPProvider(
                config("compatible"), transport=httpx.MockTransport(disconnected)
            )
        )
    assert network.value.code == "provider_network" and "PRIVATE" not in str(
        network.value
    )


def test_module_inventory_uses_canonical_items_and_indexes_item_locators(setup):
    client, app, _workspace, _provider, _fake, _secrets, canvas = setup
    course = client.post(
        "/api/workspaces", json={"kind": "canvas_course", "canvas_course_id": "101"}
    ).json()
    calls = []

    def call(method, **args):
        calls.append((method, args))
        if method == "get_course":
            return {"syllabus_body": "<p>Synthetic syllabus</p>"}
        if method == "list_modules":
            return [
                {
                    "id": 1,
                    "name": "Module",
                    "items_count": 2,
                    "items": [{"id": 999, "title": "INCOMPLETE INLINE ARRAY"}],
                }
            ]
        if method == "activity_list":
            return {
                "items": [
                    {
                        "id": 2,
                        "type": "Page",
                        "page_url": "evidence",
                        "title": "Evidence page",
                    },
                    {
                        "id": 3,
                        "type": "ExternalUrl",
                        "title": "Video reference",
                        "external_url": "https://example.com/video",
                    },
                ],
                "truncated": False,
            }
        if method == "activity_get":
            return {"id": 1, "name": "Current module"}
        return []

    canvas.client_call.side_effect = call
    base = f"/api/workspaces/{course['id']}/sources"
    inventory = client.post(base + "/inventory").json()["sources"]
    module = next(s for s in inventory if s["kind"] == "module")
    assert len(module["metadata"]["items"]) == 2 and "INCOMPLETE" not in json.dumps(
        inventory
    )
    assert any(s["source_key"] == "page:evidence" for s in inventory)
    assert (
        next(args for method, args in calls if method == "list_modules")[
            "include_items"
        ]
        is False
    )
    synced = client.post(base + "/sync", json={"source_ids": [module["id"]]}).json()[
        "sources"
    ][0]
    assert synced["status"] == "ready" and synced["title"] == "Current module"
    chunks = app.state.study_store.chunks(course["id"], module["id"])
    assert any(
        c["locator"].get("item_id") == "2" and c["locator"].get("module_id") == "1"
        for c in chunks
    )


def test_connect_native_provider_is_idempotent_and_never_stores_keys(
    setup, monkeypatch
):
    client, _app, _, _, _, secrets, _ = setup

    async def inspect(kind):
        return {"model": "synthetic-model"}

    monkeypatch.setattr("webapp.routes.inspect_login", inspect)
    first = client.post("/api/providers/local/codex")
    again = client.post("/api/providers/local/codex")
    assert first.status_code == again.status_code == 201
    assert first.json()["id"] == again.json()["id"]
    assert first.json()["credential_store"] == "native_cli"
    assert first.json()["has_key"] is False
    assert first.json()["capabilities"]["native_tools"] is False
    assert first.json()["capabilities"]["structured_output"] is True
    assert secrets.values == {}


@pytest.mark.parametrize(
    "field,value",
    [
        ("api_key", "SYNTHETIC-KEY"),
        ("remove_key", True),
        ("native_tools", True),
        ("base_url", "https://example.test/v1"),
    ],
)
def test_native_provider_rejects_api_credentials_endpoints_and_tools(
    setup, monkeypatch, field, value
):
    client, _, _, _, _, secrets, _ = setup
    monkeypatch.setattr("webapp.routes.find_cli", lambda _: "/synthetic/codex")
    result = client.post(
        "/api/providers",
        json={
            "name": "Synthetic Codex",
            "kind": "codex",
            "model": "default",
            field: value,
        },
    )
    assert result.status_code == 422
    assert secrets.values == {}


def test_native_connect_requires_web_csrf_before_inspecting_login(setup, monkeypatch):
    client, *_ = setup
    inspect = Mock()
    monkeypatch.setattr("webapp.routes.inspect_login", inspect)
    client.headers.pop("X-Workbench-CSRF")
    response = client.post("/api/providers/local/claude_code")
    assert response.status_code == 403
    inspect.assert_not_called()


def test_native_missing_login_is_actionable_without_account_details(setup, monkeypatch):
    client, *_ = setup

    async def inspect(kind):
        raise ProviderError(
            "cli_auth", "Sign in to Claude Code CLI, then connect again.", 401
        )

    monkeypatch.setattr("webapp.routes.inspect_login", inspect)
    result = client.post("/api/providers/local/claude_code")
    assert result.status_code == 401
    assert result.json()["error"]["code"] == "cli_auth"
