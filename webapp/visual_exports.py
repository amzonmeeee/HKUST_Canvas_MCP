"""Portable visual materials. Generated text is escaped, never executed."""

import re
import unicodedata
from html import escape
from io import BytesIO

from .providers import ProviderError


def wrap(value, width):
    """Wrap by display width, including Chinese text without ASCII spaces."""
    lines, line, size = [], "", 0
    for char in str(value):
        weight = 2 if unicodedata.east_asian_width(char) in "WF" else 1
        if char == "\n" or size + weight > width:
            lines.append(line.rstrip())
            line, size = "", 0
            if char == "\n":
                continue
        line += char
        size += weight
    if line:
        lines.append(line.rstrip())
    return lines or [""]


def _text(parts, x, y, value, *, width, size=20, color="#11191f", bold=False):
    lines = wrap(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value)), width)
    for index, line in enumerate(lines):
        parts.append(
            f'<text x="{x}" y="{y + index * size * 1.4}" font-size="{size}" fill="{color}" font-weight="{650 if bold else 400}">{escape(line)}</text>'
        )
    return y + len(lines) * size * 1.4


def _references(artifact):
    citations = artifact["provenance"].get("citations", [])
    numbers = {c["chunk_id"]: i + 1 for i, c in enumerate(citations)}
    return citations, lambda ids: " ".join(
        f"[{numbers[c]}]" for c in ids if c in numbers
    )


def svg_export(artifact):
    kind, content = artifact["kind"], artifact["content"]
    if kind not in {"mindmap", "infographic"}:
        raise ProviderError(
            "unsupported_export",
            "SVG export is available for maps and infographics.",
            422,
        )
    citations, refs = _references(artifact)
    parts = []
    if kind == "mindmap":
        nodes = content["nodes"]
        parents = {n["id"]: n["parent_id"] for n in nodes}
        positions, depth = {}, {}
        for node in nodes:
            current, level = node["parent_id"], 1
            while current is not None:
                level += 1
                current = parents[current]
            depth[node["id"]] = level
        width = 360 * (max(depth.values(), default=1) + 1) + 50
        y = 160
        for node in nodes:
            x = 40 + depth[node["id"]] * 360
            label_height = len(wrap(node["label"], 25)) * 28
            body_height = len(wrap(node["body"], 34)) * 22.4
            positions[node["id"]] = (x, y, label_height + body_height + 80)
            y += positions[node["id"]][2] + 22
        root_height = max(120, len(wrap(artifact["title"], 23)) * 28 + 50)
        height = max(480, y, root_height + 100)
        root_y = height / 2
        for node in nodes:
            x, ny, h = positions[node["id"]]
            px, py, ph = (
                positions[node["parent_id"]]
                if node["parent_id"]
                else (40, root_y - root_height / 2, root_height)
            )
            parts.append(
                f'<path d="M{px + 300} {py + ph / 2} H{x - 20} V{ny + h / 2} H{x}" fill="none" stroke="#005a9c" stroke-width="2"/>'
            )
        parts.append(
            f'<rect x="40" y="{root_y - root_height / 2}" width="300" height="{root_height}" rx="12" fill="#005a9c"/>'
        )
        _text(
            parts,
            60,
            root_y - root_height / 2 + 35,
            artifact["title"],
            width=23,
            size=20,
            color="#ffffff",
            bold=True,
        )
        for node in nodes:
            x, ny, h = positions[node["id"]]
            parts.append(
                f'<rect x="{x}" y="{ny}" width="300" height="{h}" rx="10" fill="#ffffff" stroke="#b2c3ce"/>'
            )
            next_y = _text(
                parts,
                x + 18,
                ny + 32,
                node["label"],
                width=25,
                size=20,
                color="#005a9c",
                bold=True,
            )
            next_y = _text(parts, x + 18, next_y + 10, node["body"], width=34, size=16)
            _text(
                parts,
                x + 18,
                next_y + 12,
                refs(node["citations"]),
                width=34,
                size=14,
                color="#005a9c",
            )
        footer_y = height + 40
    else:
        orientation = artifact["provenance"].get("orientation", "landscape")
        width, columns = {
            "landscape": (1400, 3),
            "portrait": (900, 1),
            "square": (1100, 2),
        }[orientation]
        style = artifact["provenance"].get("visual_style", "editorial")
        accent = "#00477c" if style == "bold" else "#005a9c"
        title_y = (
            _text(
                parts,
                50,
                70,
                artifact["title"],
                width=int((width - 100) / 18),
                size=32,
                color=accent,
                bold=True,
            )
            + 30
        )
        sections = content["sections"]
        card_width = (width - 100 - 24 * (columns - 1)) / columns
        body_chars = max(20, int((card_width - 40) / 10))
        y = title_y
        for offset in range(0, len(sections), columns):
            row = sections[offset : offset + columns]
            card_height = max(
                150
                + len(wrap(item["heading"], max(15, int((card_width - 40) / 12)))) * 31
                + len(wrap(item["body"], body_chars)) * 25.2
                + len(wrap(item["stat"], max(12, int((card_width - 40) / 16)))) * 38
                for item in row
            )
            for index, item in enumerate(row):
                x = 50 + index * (card_width + 24)
                radius = 20 if style == "playful" else 4
                parts.append(
                    f'<rect x="{x}" y="{y}" width="{card_width}" height="{card_height}" rx="{radius}" fill="#ffffff" stroke="#b2c3ce"/>'
                )
                yy = _text(
                    parts,
                    x + 20,
                    y + 40,
                    item["heading"],
                    width=max(15, int((card_width - 40) / 12)),
                    size=22,
                    bold=True,
                    color=accent,
                )
                if item["stat"]:
                    yy = _text(
                        parts,
                        x + 20,
                        yy + 20,
                        item["stat"],
                        width=max(12, int((card_width - 40) / 16)),
                        size=27,
                        bold=True,
                        color=accent,
                    )
                yy = _text(
                    parts, x + 20, yy + 20, item["body"], width=body_chars, size=18
                )
                _text(
                    parts,
                    x + 20,
                    yy + 18,
                    refs(item["citations"]),
                    width=body_chars,
                    size=15,
                    color=accent,
                )
                if style == "notebook":
                    parts.append(
                        f'<path d="M{x + 20} {y + card_height - 14} H{x + card_width - 20}" stroke="#b2c3ce"/>'
                    )
            y += card_height + 24
        footer_y = max(
            y + 30,
            width
            if orientation == "square"
            else width * 1.4
            if orientation == "portrait"
            else 0,
        )
    _text(parts, 50, footer_y, "Sources", width=80, size=20, bold=True)
    footer_y += 36
    for index, citation in enumerate(citations, 1):
        footer_y = (
            _text(
                parts,
                50,
                footer_y,
                f"[{index}] {citation['title']} · {citation.get('canvas_url') or ''}",
                width=int((width - 100) / 8),
                size=14,
            )
            + 8
        )
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {footer_y + 40}" role="img"><title>{escape(artifact["title"])}</title><rect width="100%" height="100%" fill="#e7ecef"/><g font-family="Arial, Microsoft JhengHei, sans-serif">'
        + "".join(parts)
        + "</g></svg>"
    ).encode()


def presentation_export(artifact):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    if artifact["kind"] != "slides":
        raise ProviderError(
            "unsupported_export", "PowerPoint export is available for slides.", 422
        )
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.author = "Canvas Workbench"
    presentation.core_properties.title = artifact["title"]
    citations, refs = _references(artifact)

    def box(slide, text, top, height, *, size=22, bold=False, color="11191F"):
        frame = slide.shapes.add_textbox(
            Inches(0.65), Inches(top), Inches(12), Inches(height)
        ).text_frame
        frame.word_wrap = True
        frame.text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        for paragraph in frame.paragraphs:
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(10)
        return frame

    for index, item in enumerate(artifact["content"]["slides"], 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("E7ECEF")
        title_size = 34
        while (
            title_size > 18
            and len(wrap(item["heading"], int(1180 / title_size))) * title_size * 1.35
            > 75
        ):
            title_size -= 1
        box(
            slide, item["heading"], 0.4, 1.1, size=title_size, bold=True, color="005A9C"
        )
        value = (
            item["body"]
            + ("\n\n" if item["body"] else "")
            + "\n".join("• " + b for b in item["bullets"])
        )
        size = 23
        while size > 12:
            lines = sum(
                len(wrap(line, int(1180 / size))) for line in value.splitlines()
            )
            if lines * size * 1.35 + len(value.splitlines()) * 10 < 335:
                break
            size -= 1
        box(slide, value, 1.65, 4.8, size=size)
        box(
            slide,
            f"{index} / {len(artifact['content']['slides'])}    Sources: {refs(item['citations'])}",
            6.65,
            0.4,
            size=12,
            color="005A9C",
        )
        slide.notes_slide.notes_text_frame.text = (
            item["notes"]
            + "\n\n"
            + "\n".join(
                f"[{i + 1}] {c['title']} {c.get('canvas_url') or ''}"
                for i, c in enumerate(citations)
                if c["chunk_id"] in item["citations"]
            )
        )
    for offset in range(0, len(citations), 5):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box(slide, "Sources", 0.4, 1, size=34, bold=True, color="005A9C")
        box(
            slide,
            "\n\n".join(
                f"[{i + offset + 1}] {c['title']}\n{c.get('canvas_url') or ''}"
                for i, c in enumerate(citations[offset : offset + 5])
            ),
            1.6,
            5.4,
            size=16,
        )
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def visual_export(artifact, format):
    return svg_export(artifact) if format == "svg" else presentation_export(artifact)
