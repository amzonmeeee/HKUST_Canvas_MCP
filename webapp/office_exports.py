"""Real Office exports with citations and portable, non-executable content."""

import re
from io import BytesIO

TEMPLATES = {
    "study_notes": "Study notes · Easy reading",
    "revision_outline": "Revision outline · Compact",
    "analysis_report": "Analysis report · Professional",
}


def _plain(value):
    # XML 1.0 cannot represent these controls, even in otherwise valid model JSON.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def word_export(artifact, template="automatic"):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    selected = (
        template
        if template != "automatic"
        else artifact["content"].get("template", "study_notes")
    )
    document = Document()
    document.core_properties.author = "Canvas Workbench"
    document.core_properties.last_modified_by = "Canvas Workbench"
    document.core_properties.title = _plain(artifact["title"])
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Cambria" if selected == "analysis_report" else "Calibri"
    normal.font.size = Pt(10 if selected == "revision_outline" else 11)
    if selected == "analysis_report":
        normal.paragraph_format.alignment = 3
    normal.paragraph_format.space_after = Pt(5 if selected == "revision_outline" else 9)
    normal.paragraph_format.line_spacing = (
        1.05 if selected == "revision_outline" else 1.2
    )
    for name in ("Normal", "Heading 1", "Heading 2", "Title"):
        document.styles[name].element.get_or_add_rPr().rFonts.set(
            qn("w:eastAsia"), "Microsoft JhengHei"
        )
    for name in ("Heading 1", "Heading 2", "Title"):
        document.styles[name].font.color.rgb = RGBColor.from_string("005A9C")
        document.styles[name].paragraph_format.keep_with_next = True
    document.add_heading(_plain(artifact["title"]), 0)
    document.add_paragraph(
        TEMPLATES.get(selected, TEMPLATES["study_notes"]), style="Subtitle"
    )
    citations = artifact["provenance"].get("citations", [])
    numbers = {c["chunk_id"]: i + 1 for i, c in enumerate(citations)}

    def refs(ids):
        return " ".join(f"[{numbers[c]}]" for c in ids if c in numbers)

    def body(text):
        for line in _plain(text).splitlines():
            if not line.strip():
                continue
            match = re.match(r"^\s*(#{1,4})\s+(.+)", line)
            if match:
                document.add_heading(match[2], min(2, len(match[1])))
            else:
                bullet = re.match(r"^\s*[-*+]\s+(.+)", line)
                numbered = re.match(r"^\s*\d+[.)]\s+(.+)", line)
                paragraph = document.add_paragraph(
                    style="List Bullet"
                    if bullet
                    else "List Number"
                    if numbered
                    else "Normal",
                )
                value = (bullet or numbered)[1] if bullet or numbered else line
                for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", value):
                    if part.startswith("**") and part.endswith("**"):
                        paragraph.add_run(part[2:-2]).bold = True
                    elif part.startswith("*") and part.endswith("*"):
                        paragraph.add_run(part[1:-1]).italic = True
                    elif part.startswith("`") and part.endswith("`"):
                        paragraph.add_run(part[1:-1]).font.name = "Consolas"
                    else:
                        paragraph.add_run(part)

    content = artifact["content"]
    kind = artifact["kind"]
    if kind in {"study_guide", "document"}:
        for item in content["sections"]:
            document.add_heading(_plain(item["heading"]), 1)
            body(item["body"])
            document.add_paragraph(refs(item["citations"]), style="Caption")
    elif kind == "spreadsheet":
        table = document.add_table(rows=1, cols=len(content["columns"]))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, content["columns"]):
            cell.text = _plain(value)
        for row in content["rows"]:
            for cell, value in zip(table.add_row().cells, row["cells"]):
                cell.text = _plain(value)
            document.add_paragraph(refs(row["citations"]), style="Caption")
    elif kind == "quiz":
        for index, item in enumerate(content["questions"], 1):
            document.add_heading(f"{index}. {_plain(item['question'])}", 1)
            for index, choice in enumerate(item["choices"]):
                document.add_paragraph(f"{chr(65 + index)}. {_plain(choice)}")
            document.add_paragraph(f"Answer: {chr(65 + item['answer_index'])}")
            body(item["explanation"])
            document.add_paragraph(refs(item["citations"]), style="Caption")
    else:
        for item in content["cards"]:
            document.add_heading(_plain(item["front"]), 1)
            body(item["back"])
            document.add_paragraph(refs(item["citations"]), style="Caption")
    document.add_heading("Sources", 1)
    for index, citation in enumerate(citations, 1):
        document.add_paragraph(
            f"[{index}] {_plain(citation['title'])}\n{_plain(citation.get('canvas_url') or '')}\n{_plain(citation['excerpt'])}"
        )
    document.add_heading("Generation details", 1)
    provenance = artifact["provenance"]
    for label, key in [
        ("Provider", "provider_name"),
        ("Model", "model"),
        ("Created", "created_at"),
        ("Instructions", "prompt"),
    ]:
        if provenance.get(key):
            document.add_paragraph(f"{label}: {_plain(provenance[key])}")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Canvas Workbench · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def excel_export(artifact):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.properties.creator = "Canvas Workbench"
    workbook.properties.lastModifiedBy = "Canvas Workbench"
    workbook.properties.title = _plain(artifact["title"])
    sheet = workbook.active
    sheet.title = "Study material"
    content = artifact["content"]
    kind = artifact["kind"]
    citations = artifact["provenance"].get("citations", [])
    numbers = {c["chunk_id"]: str(i + 1) for i, c in enumerate(citations)}

    def refs(ids):
        return ", ".join(numbers[c] for c in ids if c in numbers)

    if kind == "spreadsheet":
        rows = [content["columns"] + ["Sources"]] + [
            row["cells"] + [refs(row["citations"])] for row in content["rows"]
        ]
    elif kind in {"document", "study_guide"}:
        rows = [["Heading", "Notes", "Sources"]] + [
            [row["heading"], row["body"], refs(row["citations"])]
            for row in content["sections"]
        ]
    elif kind == "flashcards":
        rows = [["Prompt", "Answer", "Sources"]] + [
            [row["front"], row["back"], refs(row["citations"])]
            for row in content["cards"]
        ]
    else:
        rows = [
            ["Question", "A", "B", "C", "D", "Answer", "Explanation", "Sources"]
        ] + [
            [
                row["question"],
                *row["choices"],
                chr(65 + row["answer_index"]),
                row["explanation"],
                refs(row["citations"]),
            ]
            for row in content["questions"]
        ]

    def fill(worksheet, values):
        for r, row in enumerate(values, 1):
            for c, value in enumerate(row, 1):
                cell = worksheet.cell(r, c, _plain(value)[:32767])
                cell.data_type = "s"  # Never interpret model/source text as a formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(
                    name="Calibri",
                    size=11,
                    color="FFFFFF" if r == 1 else "11191F",
                    bold=r == 1,
                )
                if r == 1 or r % 2 == 0:
                    cell.fill = PatternFill(
                        "solid", fgColor="005A9C" if r == 1 else "E7ECEF"
                    )
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 28
        for c in range(1, len(values[0]) + 1):
            worksheet.column_dimensions[get_column_letter(c)].width = min(
                60,
                max(
                    18,
                    max(len(str(row[c - 1])) for row in values if c <= len(row)) * 0.8,
                ),
            )
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_title_rows = "1:1"

    fill(sheet, rows)
    fill(
        workbook.create_sheet("Sources"),
        [["Source", "Title", "Canvas URL", "Excerpt"]]
        + [
            [i + 1, c["title"], c.get("canvas_url") or "", c["excerpt"]]
            for i, c in enumerate(citations)
        ],
    )
    fill(
        workbook.create_sheet("Generation details"),
        [["Field", "Value"]]
        + [
            [key, artifact["provenance"].get(key) or ""]
            for key in (
                "provider_name",
                "model",
                "created_at",
                "difficulty",
                "topic",
                "prompt",
            )
        ],
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
