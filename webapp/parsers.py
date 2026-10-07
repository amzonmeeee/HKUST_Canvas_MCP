"""Bounded, offline text extraction with stable document locators."""

from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from html.parser import HTMLParser
from pathlib import Path

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
SUPPORTED = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".pptx", ".vtt", ".srt"}


class ParseError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.heading = ""
        self.heading_text = []
        self.in_heading = False
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "iframe", "object", "svg", "template"}:
            self.hidden += 1
        if self.hidden:
            return
        if re.fullmatch("h[1-6]", tag):
            self.in_heading = True
            self.heading_text = []
        if tag in {"p", "br", "div", "li", "tr", "section", "article"}:
            self.parts.append(("\n", self.heading))

    def handle_endtag(self, tag):
        if tag in {"script", "style", "iframe", "object", "svg", "template"}:
            self.hidden = max(0, self.hidden - 1)
        if self.hidden:
            return
        if re.fullmatch("h[1-6]", tag):
            self.heading = "".join(self.heading_text).strip()[:200]
            self.in_heading = False
            self.parts.append(("\n" + self.heading + "\n", self.heading))
        if tag in {"p", "div", "li", "tr", "section", "article"}:
            self.parts.append(("\n", self.heading))

    def handle_data(self, data):
        if not self.hidden:
            if self.in_heading:
                self.heading_text.append(data)
            else:
                self.parts.append((data, self.heading))


def html_sections(text):
    parser = HTMLText()
    parser.feed(text[:MAX_TEXT_CHARS])
    sections = []
    for part, heading in parser.parts:
        if sections and sections[-1][1].get("heading") == heading:
            sections[-1] = (sections[-1][0] + part, sections[-1][1])
        else:
            sections.append((part, {"heading": heading} if heading else {}))
    return sections


def chunks_for_sections(sections):
    chunks, total = [], 0
    for text, locator in sections:
        text = re.sub(r"[ \t]+", " ", text).strip()
        if not text:
            continue
        total += len(text)
        if total > MAX_TEXT_CHARS:
            raise ParseError(
                "text_too_large",
                "Extracted text exceeds the 2 million character limit.",
            )
        offset = 0
        while offset < len(text):
            end = min(offset + 1800, len(text))
            if end < len(text):
                boundary = max(
                    text.rfind("\n", offset + 900, end),
                    text.rfind(" ", offset + 900, end),
                )
                if boundary > offset:
                    end = boundary
            fragment = text[offset:end].strip()
            if fragment:
                chunks.append(
                    {"text": fragment, "locator": {**locator, "offset": offset}}
                )
            if end == len(text):
                break
            offset = max(end - 160, offset + 1)
    if not chunks:
        raise ParseError(
            "no_extractable_text",
            "No readable text was found. Scanned PDFs need OCR outside this app.",
        )
    return chunks


def check_zip(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 10000
            or sum(i.file_size for i in entries) > 100 * 1024 * 1024
        ):
            raise ParseError(
                "document_too_large", "The document expands beyond the supported limit."
            )
        if any(i.flag_bits & 1 for i in entries):
            raise ParseError(
                "encrypted_document", "Password-protected documents cannot be indexed."
            )


def extract(path: Path, suffix: str):
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ParseError("file_too_large", "Files must be at most 25 MB.")
    if suffix not in SUPPORTED:
        raise ParseError(
            "unsupported_format",
            "This file type is kept as a reference; text extraction is unavailable.",
        )
    if suffix in {".docx", ".pptx"}:
        check_zip(path)
    sections = []
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise ParseError(
                "encrypted_document", "Password-protected PDFs cannot be indexed."
            )
        if len(reader.pages) > 1000:
            raise ParseError("document_too_large", "PDFs are limited to 1,000 pages.")
        total = 0
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise ParseError("text_too_large", "The PDF contains too much text.")
            sections.append((text, {"page": index + 1}))
    elif suffix == ".docx":
        from docx import Document

        document = Document(path)
        heading = ""
        for index, paragraph in enumerate(document.paragraphs):
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                heading = paragraph.text[:200]
            sections.append(
                (paragraph.text, {"paragraph": index + 1, "heading": heading})
            )
        for index, table in enumerate(document.tables):
            sections.append(
                (
                    "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows),
                    {"table": index + 1},
                )
            )
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        for index, slide in enumerate(presentation.slides):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(
                        " | ".join(c.text for c in row.cells)
                        for row in shape.table.rows
                    )
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            sections.append(("\n".join(parts), {"slide": index + 1}))
    else:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ParseError(
                "text_encoding", "Text uploads must use UTF-8 encoding."
            ) from exc
        if suffix in {".html", ".htm"}:
            sections = html_sections(text)
        elif suffix in {".vtt", ".srt"}:
            for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
                lines = block.splitlines()
                timed = next((i for i, line in enumerate(lines) if "-->" in line), None)
                if timed is not None:
                    timestamp = lines[timed].split("-->")[0].strip().replace(",", ".")
                    sections.append(
                        (
                            re.sub("<[^>]+>", "", "\n".join(lines[timed + 1 :])),
                            {"timestamp": timestamp},
                        )
                    )
        elif suffix == ".md":
            heading, parts = "", []
            for line in text.splitlines():
                if re.match(r"^#{1,6} ", line):
                    sections.append(("\n".join(parts), {"heading": heading}))
                    heading, parts = line.lstrip("# ")[:200], []
                parts.append(line)
            sections.append(("\n".join(parts), {"heading": heading}))
        else:
            sections = [(text, {})]
    return chunks_for_sections(sections)


def parse_file(path: Path, suffix: str):
    """Isolate document parsers so malformed documents cannot stall the web server."""
    try:
        result = subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), str(path), suffix],
            capture_output=True,
            text=True,
            timeout=45,
            cwd=path.parent,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise ParseError(
            "parser_timeout", "Extraction took too long. Try a smaller document."
        ) from exc
    try:
        payload = json.loads(result.stdout)
    except (ValueError, TypeError) as exc:
        raise ParseError("parser_failed", "The document could not be parsed.") from exc
    if "error" in payload:
        raise ParseError(payload["error"]["code"], payload["error"]["message"])
    return payload["chunks"]


if __name__ == "__main__":
    try:
        if sys.platform != "win32":
            import resource

            resource.setrlimit(resource.RLIMIT_CPU, (35, 35))
            # macOS reserves a large Python heap arena: RLIMIT_DATA rejects even
            # tiny text files there. Keep CPU/timeout bounds on macOS, add an
            # address-space bound where the kernel supports it predictably.
            if sys.platform == "linux":
                resource.setrlimit(
                    resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024)
                )
        print(json.dumps({"chunks": extract(Path(sys.argv[1]), sys.argv[2])}))
    except ParseError as exc:
        print(json.dumps({"error": {"code": exc.code, "message": str(exc)}}))
    except Exception:  # noqa: BLE001 -- untrusted parser errors may contain private text
        print(
            json.dumps(
                {
                    "error": {
                        "code": "parser_failed",
                        "message": "The document is malformed or unsupported.",
                    }
                }
            )
        )
