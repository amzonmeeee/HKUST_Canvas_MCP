"""Owned desktop backend lifecycle. No browser/Keychain data is bundled or copied."""

from __future__ import annotations

import argparse
import json
import os
import re
import signal
import sys
import tempfile
import time
from pathlib import Path
from urllib.error import URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen

from .db import default_data_dir, validate_data_dir


def write_private_json(path: Path, payload: dict):
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    descriptor, name = tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(descriptor, "w") as stream:
            json.dump(payload, stream)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def verified_session(marker: Path) -> str | None:
    try:
        payload = json.loads(marker.read_text())
        url = payload["url"]
        parsed = urlsplit(url)
        if (
            parsed.scheme != "http"
            or parsed.hostname != "127.0.0.1"
            or parsed.username
            or parsed.password
            or parsed.path != "/"
            or parsed.query
            or not parsed.port
        ):
            return None
        if not re.fullmatch(r"session=[A-Za-z0-9_-]{32,128}", parsed.fragment):
            return None
        base = f"http://127.0.0.1:{parsed.port}"
        req = Request(
            base + "/api/session",
            data=b"",
            headers={
                "Origin": base,
                "Authorization": "Bearer " + parsed.fragment.split("=", 1)[1],
            },
            method="POST",
        )
        with urlopen(req, timeout=1) as response:
            if response.status == 200 and isinstance(
                json.loads(response.read(4096)).get("csrf_token"), str
            ):
                return url
    except (OSError, ValueError, KeyError, TypeError, URLError):
        pass
    return None


def run_desktop(*, data_dir: Path, ready_file: Path, port=8765, parent_pid=None):
    import fcntl

    from .launcher import launch

    data_dir = validate_data_dir(data_dir)
    state = data_dir / "launcher"
    state.mkdir(parents=True, mode=0o700, exist_ok=True)
    os.chmod(state, 0o700)
    marker = state / "session.json"
    descriptor = os.open(state / "server.lock", os.O_CREAT | os.O_RDWR, 0o600)
    try:
        deadline = time.monotonic() + 30
        while True:
            try:
                fcntl.flock(descriptor, fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except BlockingIOError:
                existing = verified_session(marker)
                if existing:
                    write_private_json(ready_file, {"url": existing, "owned": False})
                    return
                if time.monotonic() >= deadline:
                    raise RuntimeError(
                        "The existing workbench is not responding. Quit its launcher and retry."
                    )
                time.sleep(0.1)

        def ready(url):
            write_private_json(marker, {"url": url, "pid": os.getpid()})
            write_private_json(ready_file, {"url": url, "owned": True})

        # Uvicorn re-emits SIGTERM after graceful shutdown. Keep control long
        # enough to remove private session metadata and release our lock.
        previous_term = signal.signal(signal.SIGTERM, lambda *_: None)
        try:
            launch(
                port=port,
                open_browser=False,
                data_dir=data_dir,
                on_ready=ready,
                parent_pid=parent_pid,
            )
        finally:
            marker.unlink(missing_ok=True)
            signal.signal(signal.SIGTERM, previous_term)
    finally:
        os.close(descriptor)


def self_test():
    import asyncio

    from docx import Document
    from fastapi.testclient import TestClient
    from pptx import Presentation
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    from canvas_mcp.server import WRITE_TOOL_NAMES, mcp

    from .api import create_app
    from .db import SCHEMA_VERSION
    from .parsers import parse_file

    with tempfile.TemporaryDirectory(prefix="workbench-package-smoke-") as folder:
        root = Path(folder)
        document = Document()
        document.add_paragraph("Synthetic package evidence")
        document.save(root / "example.docx")
        slides = Presentation()
        slides.slides.add_slide(
            slides.slide_layouts[1]
        ).shapes.title.text = "Synthetic package evidence"
        slides.save(root / "example.pptx")
        (root / "example.html").write_text("<p>Synthetic package evidence</p>")
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (Synthetic package evidence) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(root / "example.pdf")
        for extension in (".pdf", ".docx", ".pptx", ".html"):
            assert any(
                "Synthetic package evidence" in chunk["text"]
                for chunk in parse_file(root / ("example" + extension), extension)
            )
        origin = "http://127.0.0.1:8765"
        with TestClient(
            create_app(
                data_dir=Path(folder) / "data", launch_secret="synthetic-package"
            ),
            base_url=origin,
        ) as client:
            assert client.get("/").status_code == 200
            result = client.post(
                "/api/session",
                headers={"Origin": origin, "Authorization": "Bearer synthetic-package"},
            )
            client.headers.update(
                {"Origin": origin, "X-Workbench-CSRF": result.json()["csrf_token"]}
            )
            assert client.get("/api/onboarding").json()["completed"] is False
            assert (
                client.get("/api/canvas/status").json()["auth_status"] == "unconfigured"
            )
            assert client.get("/api/canvas/courses").status_code == 409
            visible = {tool.name for tool in asyncio.run(mcp.list_tools())}
            assert not visible.intersection(WRITE_TOOL_NAMES)
            assert (
                client.put(
                    "/api/onboarding", json={"privacy_acknowledged": True}
                ).status_code
                == 200
            )
            workspace = client.post(
                "/api/workspaces", json={"title": "Synthetic package workspace"}
            )
            assert workspace.status_code == 201
            source = client.post(
                f"/api/workspaces/{workspace.json()['id']}/sources/text",
                json={
                    "title": "Synthetic source",
                    "content": "Package extraction must run in its isolated worker.",
                },
            )
            assert source.status_code == 201 and source.json()["status"] == "ready"
    print(
        json.dumps({"self_test": "passed", "schema": SCHEMA_VERSION, "frontend": True})
    )


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv[:1] == ["--parse-worker"]:
        from .parsers import worker_main

        return worker_main(argv[1:])
    if argv[:1] == ["--mcp"]:
        from mcp_entry import main as mcp_main

        return mcp_main(argv[1:])
    if argv[:1] == ["--cli"]:
        from canvas_cli import main as cli_main

        sys.argv = [sys.argv[0], *argv[1:]]
        return cli_main()
    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--data-dir", type=Path, default=default_data_dir())
    parser.add_argument("--ready-file", type=Path)
    parser.add_argument("--parent-pid", type=int)
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args(argv)
    if args.self_test:
        return self_test()
    if not args.ready_file or not 0 <= args.port <= 65535:
        parser.error("A private --ready-file and a port from 0 to 65535 are required.")
    run_desktop(
        data_dir=args.data_dir,
        ready_file=args.ready_file,
        port=args.port,
        parent_pid=args.parent_pid,
    )


if __name__ == "__main__":
    main()
